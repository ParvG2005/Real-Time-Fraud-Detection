"""Build the architecture graphic and an editable 14-slide presentation."""

from pathlib import Path
import html
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
FONT = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
font = lambda n: (
    ImageFont.truetype(str(FONT), n)
    if FONT.exists()
    else ImageFont.load_default(size=n)
)
image = Image.new("RGB", (1600, 960), "#09111f")
draw = ImageDraw.Draw(image)
draw.text(
    (70, 45), "FraudShield AI / Local architecture", font=font(43), fill="#cde4ff"
)
draw.text(
    (72, 108),
    "A reproducible decision pipeline with persistent evidence and human oversight.",
    font=font(23),
    fill="#7e9fcb",
)
boxes = {
    "ui": (
        70,
        210,
        390,
        340,
        "React dashboard",
        "Transactions · alerts · investigations",
    ),
    "api": (
        535,
        210,
        990,
        340,
        "FastAPI API",
        "JWT / RBAC · validation · orchestration",
    ),
    "ml": (
        1130,
        210,
        1530,
        340,
        "FastAPI ML service",
        "XGBoost · Isolation Forest · TreeSHAP",
    ),
    "db": (
        70,
        465,
        450,
        595,
        "PostgreSQL + pgvector",
        "Transactions · labels · cosine retrieval",
    ),
    "policy": (
        555,
        465,
        990,
        595,
        "Hybrid risk policy",
        "35% ML · 20% rules · 25% behavior · 10% anomaly · 10% history",
    ),
    "redis": (
        1130,
        465,
        1530,
        595,
        "Redis sliding windows",
        "Exact velocity · rate limits · fallback",
    ),
    "feedback": (
        70,
        725,
        450,
        855,
        "Human investigation",
        "Confirm · clear · escalate · feedback",
    ),
    "decision": (
        555,
        725,
        990,
        855,
        "ALLOW / REVIEW / BLOCK",
        "Commit evidence → publish live event",
    ),
    "explain": (
        1130,
        725,
        1530,
        855,
        "Asynchronous explanation",
        "Local evidence template · optional Bedrock",
    ),
}


def arrow(start, end, label=None):
    draw.line([start, end], fill="#4c91c7", width=4)
    x, y = end
    dx = x - start[0]
    dy = y - start[1]
    if abs(dx) > abs(dy):
        points = [
            (x, y),
            (x - (12 if dx > 0 else -12), y - 7),
            (x - (12 if dx > 0 else -12), y + 7),
        ]
    else:
        points = [
            (x, y),
            (x - 7, y - (12 if dy > 0 else -12)),
            (x + 7, y - (12 if dy > 0 else -12)),
        ]
    draw.polygon(points, fill="#4c91c7")
    if label:
        draw.text(
            ((start[0] + x) // 2 - 40, (start[1] + y) // 2 - 28),
            label,
            font=font(18),
            fill="#7e9fcb",
        )


arrow((390, 275), (535, 275), "REST / WS")
arrow((990, 275), (1130, 275), "inference")
arrow((760, 340), (760, 465))
arrow((555, 530), (450, 530))
arrow((990, 530), (1130, 530))
arrow((760, 595), (760, 725))
arrow((555, 790), (450, 790))
arrow((990, 790), (1130, 790))
arrow((260, 725), (260, 595), "feedback")
for key, (x1, y1, x2, y2, title, subtitle) in boxes.items():
    bg = "#18375a" if key in ("api", "decision") else "#13213a"
    fg = "#d8ecff" if key in ("api", "decision") else "#a8cff3"
    draw.rounded_rectangle(
        (x1, y1, x2, y2), radius=19, fill=bg, outline="#2a4867", width=2
    )
    draw.text((x1 + 23, y1 + 27), title, font=font(26), fill=fg)
    # wrap the policy detail deliberately to keep lines readable.
    words = subtitle.split()
    lines = []
    line = ""
    for word in words:
        candidate = (line + " " + word).strip()
        if draw.textlength(candidate, font=font(17)) > x2 - x1 - 45:
            lines.append(line)
            line = word
        else:
            line = candidate
    if line:
        lines.append(line)
    for i, line in enumerate(lines):
        draw.text(
            (x1 + 23, y1 + 73 + i * 20),
            line,
            font=font(17),
            fill="#7ab9e9" if key in ("api", "decision") else "#7394be",
        )
draw.text(
    (72, 903),
    "Local demo: synthetic data, real ML inference, actual database records. The LLM never determines the score.",
    font=font(20),
    fill="#7191ba",
)
image.save(DOCS / "architecture.png")
image.save(ROOT / "architecture.png")
slides = [
    (
        "FraudShield AI",
        "Real-time fraud intelligence.\nEvery signal explained.",
        [
            "Digital lending · hybrid detection · human investigation",
            "Local interview demonstration | FastAPI + React",
        ],
        None,
    ),
    (
        "The problem",
        "Static rules see fragments. Fraud spans behavior.",
        [
            "Unusual amounts may be legitimate; a familiar device may be compromised.",
            "Signals need context: history, velocity, device and location.",
            "Investigators need evidence they can inspect, not an unexplained score.",
        ],
        None,
    ),
    (
        "The solution",
        "Detect → score → explain → investigate → learn",
        [
            "Combine supervised prediction, deterministic rules and behavioral signals.",
            "Persist the features and evidence behind every policy decision.",
            "Keep the analyst in control of the investigation label.",
        ],
        None,
    ),
    (
        "Local architecture",
        "Five services. One complete workflow.",
        [],
        "architecture.png",
    ),
    (
        "The decision pipeline",
        "A transaction becomes an auditable decision.",
        [
            "Validate identity, role, amount and request schema.",
            "Derive features from committed history and Redis windows.",
            "Run ML + rules + behavior + anomaly + pgvector retrieval.",
            "Aggregate the score; commit evidence; notify the dashboard.",
        ],
        None,
    ),
    (
        "Machine learning",
        "Train and measure. Never hardcode predictions.",
        [
            "18,000 synthetic examples; 60 / 20 / 20 train-validation-test split.",
            "XGBoost model selection uses validation PR-AUC only.",
            "Isolation Forest detects departures from legitimate training behavior.",
            "No real-world fraud prevalence or deployment performance is claimed.",
        ],
        None,
    ),
    (
        "Explainability",
        "Show why the model moved.",
        [
            "TreeSHAP contributions are measured in log-odds.",
            "Contributions + base reconstruct the actual model prediction.",
            "Local prose cites captured evidence; it is explicitly not LLM output.",
        ],
        "explainability.png",
    ),
    (
        "Historical cases",
        "Similar patterns add context, not certainty.",
        [
            "pgvector executes actual cosine similarity search.",
            "256-dimensional local hashing vectors represent pattern text.",
            "Only confirmed cases from the matching embedding provider are retrieved.",
            "Reference cases are labeled synthetic; analyst confirmations enter retrieval.",
        ],
        None,
    ),
    (
        "Security and consistency",
        "Make the boundaries explicit.",
        [
            "JWT + bcrypt + ADMIN / ANALYST / VIEWER roles.",
            "Parameterized SQL, strict validation and minimized identifiers.",
            "Idempotency protects retries; per-customer locks protect velocity counts.",
            "Audit events capture access, policy edits and investigation decisions.",
        ],
        None,
    ),
    ("The dashboard", "A workspace for fraud operations.", [], "dashboard.png"),
    (
        "Live demonstration",
        "Normal → review → block → investigate",
        [
            "Known-device repayment → ALLOW.",
            "Unusual amount on familiar device → REVIEW.",
            "New device + new city + real burst → BLOCK.",
            "Confirm fraud, or mark legitimate and retain the original policy output.",
        ],
        None,
    ),
    (
        "Measured results",
        "Synthetic holdout performance — not a production claim.",
        [
            "Precision 88.83% | Recall 88.38% | F1 88.61%.",
            "PR-AUC 0.9285 | ROC-AUC 0.9779.",
            "False positive rate 1.37% | False negative rate 11.62%.",
            "The dashboard separately reports selected analyst-review feedback.",
        ],
        None,
    ),
    (
        "Responsible AI and future scope",
        "A useful prototype has honest limits.",
        [
            "Risk is not proof; human review can correct false positives.",
            "No funds are moved, no creditworthiness is assessed, no users are banned.",
            "Next: representative data, calibration, fairness audits and durable jobs.",
            "At scale: shared events, vector indexes, model registry and rollback.",
        ],
        None,
    ),
    (
        "FraudShield AI",
        "From reactive rules to explainable fraud intelligence.",
        [
            "Real-time detection. Persistent evidence. Human oversight.",
            "Run the complete demo locally: http://localhost:3000",
            "Read the code, repeat the tests, inspect every decision.",
        ],
        None,
    ),
]
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
sections = []
for index, (title, subtitle, bullets, asset) in enumerate(slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("09111F")

    def text_box(x, y, w, h, text, size, color="CDE4FF", bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Aptos"
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
        return shape

    text_box(
        0.6, 0.35, 11, 0.4, f"FRAUDSHIELD AI  /  {index + 1:02d}", 11, "76A0CF", True
    )
    text_box(0.6, 0.95, 12, 0.65, title, 30, bold=True)
    text_box(0.6, 1.72, 12, 0.8, subtitle, 24, "8EC9F5")
    if asset and (DOCS / asset).exists():
        if bullets:
            text_box(
                0.6,
                2.95,
                4.2,
                3.8,
                "\n\n".join("• " + b for b in bullets),
                16,
                "89A6CE",
            )
            picture = Image.open(DOCS / asset)
            scale = min(7.55 / picture.width, 4.05 / picture.height)
            width, height = picture.width * scale, picture.height * scale
            slide.shapes.add_picture(
                str(DOCS / asset),
                Inches(5.15 + (7.55 - width) / 2),
                Inches(2.8),
                width=Inches(width),
                height=Inches(height),
            )
        else:
            picture = Image.open(DOCS / asset)
            scale = min(11.6 / picture.width, 4.45 / picture.height)
            width, height = picture.width * scale, picture.height * scale
            slide.shapes.add_picture(
                str(DOCS / asset),
                Inches(0.85 + (11.6 - width) / 2),
                Inches(2.5),
                width=Inches(width),
                height=Inches(height),
            )
    else:
        text_box(
            0.75, 3, 11.6, 3.8, "\n\n".join("• " + b for b in bullets), 22, "8CAACE"
        )
    text_box(
        0.6,
        7.1,
        12,
        0.25,
        "LOCAL DEMONSTRATION  ·  SYNTHETIC DATA  ·  REAL INFERENCE",
        9,
        "5D81AD",
    )
    slide.notes_slide.notes_text_frame.text = (
        "Refer to docs/PRESENTATION.md for the full talk track. " + " ".join(bullets)
    )
    body = (
        "<ul>" + "".join("<li>" + html.escape(b) + "</li>" for b in bullets) + "</ul>"
        if bullets
        else ""
    )
    if asset and (DOCS / asset).exists():
        body = (
            f'<div class="content {"split" if bullets else ""}">'
            + body
            + f'<img src="{asset}"/></div>'
        )
    sections.append(
        f'<section><div class="eyebrow">FRAUDSHIELD AI / {index + 1:02d}</div><h1>{html.escape(title)}</h1><h2>{html.escape(subtitle).replace(chr(10), "<br>")}</h2>{body}<footer>LOCAL DEMONSTRATION · SYNTHETIC DATA · REAL INFERENCE <span>{index + 1:02d} / 14</span></footer></section>'
    )
prs.save(DOCS / "FraudShield-Presentation.pptx")
css = """@page{size:1600px 900px;margin:0}*{box-sizing:border-box}body{margin:0;font-family:Arial,sans-serif;background:#09111f;color:#cde4ff}section{position:relative;width:1600px;height:900px;padding:64px 80px;page-break-after:always;overflow:hidden}.eyebrow{letter-spacing:3px;font-size:18px;color:#6e9dce;font-weight:600}h1{font-size:55px;letter-spacing:-1.8px;margin:36px 0 16px}h2{font-size:34px;font-weight:400;color:#8dbfed;line-height:1.3;margin:0 0 46px}ul{padding-left:30px;max-width:1300px}li{font-size:29px;line-height:1.5;color:#8caacd;margin:28px 0}.content{height:550px;display:flex;justify-content:center}.content img{max-width:100%;max-height:100%;object-fit:contain}.split{gap:40px}.split ul{width:40%}.split li{font-size:24px}.split img{width:57%;object-position:top}footer{position:absolute;bottom:32px;left:80px;right:80px;font-size:14px;letter-spacing:2px;color:#688bb5}footer span{float:right}@media screen{body{background:#060a13}section{margin:30px auto;background:#09111f;box-shadow:0 15px 50px #1d341c20}}"""
(DOCS / "presentation.html").write_text(
    '<!doctype html><html lang="en"><head><meta charset="UTF-8"><title>FraudShield AI Presentation</title><style>'
    + css
    + "</style></head><body>"
    + "".join(sections)
    + "</body></html>"
)
print("Created architecture.png, editable PPTX and HTML slides.")
